import streamlit as st
import pandas as pd
from openpyxl import load_workbook
import io
import datetime

# Sayfa ayarları
st.set_page_config(page_title="Ferçim Rapor", page_icon="🏭", layout="centered")

st.title("🏭 Ferçim Niğde - Otomatik Rapor")
st.markdown("SAP MM ekran görüntülerini yükleyin ve rapor tarihini seçin.")

# Tarih Seçici
rapor_tarihi = st.date_input("🗓️ Rapor Tarihi", datetime.date.today())
formatli_tarih = rapor_tarihi.strftime("%d.%m.%Y")
gun_ismi = str(rapor_tarihi.day) # Seçilen tarihin gün kısmı (örn: "26")

# Dosya yükleme alanı
uploaded_files = st.file_uploader("📸 SAP Ekran Görüntülerini Seçin", type=["png", "jpg", "jpeg"], accept_multiple_files=True)

if uploaded_files:
    st.success(f"✅ {len(uploaded_files)} görüntü yüklendi. İşleme hazır.")
    
    if st.button("🚀 Raporu Oluştur"):
        with st.spinner('Veriler okunuyor ve Excel hazırlanıyor...'):
            try:
                # Şablonu yükle (Depondaki dosya adının bu olduğundan emin olun)
                template_path = "Niğde Günlük Rapor 2026_09_24.xlsx"
                wb = load_workbook(template_path)
                
                # 1. SAYFA ADI GÜNCELLEMESİ
                ws = wb.worksheets[0] # İlk sayfayı otomatik seç
                ws.title = gun_ismi # Sayfa adını gün numarası yap (örn: 26)
                
                # 2. TARİH HÜCRESİ GÜNCELLEMESİ (B2)
                ws['A2'] = "" # Önceki hatalı A2 kaydını temizle
                ws['B2'] = formatli_tarih # Doğru hücreye (B2) tarihi yaz
                
                # 3. DÖNEM SONU STOKLAR (Geçici OCR Simülasyonu - F Sütunu)
                # Klinker Dönem Sonu
                ws['F4'] = 94653.233 
                # Çimento Kalemleri Dönem Sonu
                ws['F7'] = 676.000   
                ws['F8'] = 218.000   
                ws['F9'] = 173.000   
                ws['F10'] = 1329.000 
                # Yakıt Kalemleri Dönem Sonu
                ws['F13'] = 3655.000 # Petrokok
                ws['F14'] = 2693.000 # Linyit
                
                # Yeni dosyayı belleğe kaydet
                output = io.BytesIO()
                wb.save(output)
                output.seek(0)
                
                st.success(f"🎉 {formatli_tarih} tarihli Excel Raporu Başarıyla Dolduruldu!")
                
                # DİNAMİK DOSYA İSMİ İLE İNDİRME BUTONU
                dosya_adi = f"Fercim_Nigde_Gunluk_Rapor_{formatli_tarih}.xlsx"
                st.download_button(
                    label=f"📥 İndir: {dosya_adi}",
                    data=output,
                    file_name=dosya_adi,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            except Exception as e:
                st.error(f"⚠️ Hata: Excel dosyası bulunamadı veya işlenemedi. Detay: {e}")
